"""月度课时导出：生成「高一（1）部课时（分班后）」式的工作簿（6 个页签）。

- 周一至周五课 / 周一至周五自习 / 周六日课 / 周六日自习：来自该月实际课程记录
- 详细课表：来自所选模板
- 监考：仅保留教师列
"""

from collections import defaultdict
from datetime import date, timedelta

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from timetable.models import (
    CourseSession,
    ReportGroup,
    SchoolClass,
    Teacher,
    TeachingAssignment,
    TimeSlot,
)

WEEKDAY_CN = ["一", "二", "三", "四", "五", "六", "日"]

# 各表的固定子列顺序（按 kind）
COURSE_KINDS = [ReportGroup.Kind.MORNING, ReportGroup.Kind.AFTERNOON]
STUDY_KINDS = [ReportGroup.Kind.EARLY_READ, ReportGroup.Kind.SELF_STUDY]
DETAIL_KINDS = [
    ReportGroup.Kind.MORNING,
    ReportGroup.Kind.AFTERNOON,
    ReportGroup.Kind.EARLY_READ,
    ReportGroup.Kind.SELF_STUDY,
]
KIND_LABEL = dict(ReportGroup.Kind.choices)

_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _month_bounds(year: int, month: int) -> tuple[date, date]:
    start = date(year, month, 1)
    end = (
        date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    ) - timedelta(days=1)
    return start, end


def _is_weekend(day: date) -> bool:
    return day.isoweekday() >= 6


def _date_label(day: date) -> str:
    return f"周{WEEKDAY_CN[day.isoweekday() - 1]} {day.month}.{day.day}"


def _finish(ws, nrows: int, ncols: int, freeze: str | None = "E3") -> None:
    for r in range(1, nrows + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER
    for r in (1, 2):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).font = Font(bold=True)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 9
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 20
    for c in range(5, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 7.5
    if freeze:
        ws.freeze_panes = freeze


def build_monthly_workbook(year: int, month: int, template) -> openpyxl.Workbook:
    start, end = _month_bounds(year, month)
    slot_kind = dict(TimeSlot.objects.values_list("id", "report_group__kind"))

    sessions = CourseSession.objects.filter(
        date__range=(start, end), teacher__isnull=False
    ).select_related("teacher")
    count: dict[tuple, int] = defaultdict(int)  # (tid, date, kind) -> 次数
    date_kind: set[tuple] = set()  # 该日期该分类是否存在课程
    teacher_classes: dict[int, set[int]] = defaultdict(set)
    for s in sessions:
        kind = slot_kind.get(s.time_slot_id)
        if not kind:
            continue
        count[(s.teacher_id, s.date, kind)] += 1
        date_kind.add((s.date, kind))
        teacher_classes[s.teacher_id].add(s.school_class_id)

    assign = {
        (a.school_class_id, a.subject_id): a.teacher_id
        for a in TeachingAssignment.objects.all()
    }
    tpl_count: dict[tuple, int] = defaultdict(int)  # (tid, weekday, kind) -> 次数
    tpl_classes: dict[int, set[int]] = defaultdict(set)
    for entry in template.entries.select_related("time_slot").all():
        kind = slot_kind.get(entry.time_slot_id)
        if not kind:
            continue
        tid = assign.get((entry.school_class_id, entry.subject_id))
        if not tid:
            continue
        tpl_count[(tid, entry.weekday, kind)] += 1
        tpl_classes[tid].add(entry.school_class_id)

    teacher_ids = set(teacher_classes) | set(tpl_classes)
    teachers = list(
        Teacher.objects.filter(id__in=teacher_ids, is_active=True).select_related(
            "subject"
        )
    )  # 已按 subject__sort_order, id 排序

    class_names = dict(SchoolClass.objects.values_list("id", "name"))
    class_order = dict(SchoolClass.objects.values_list("id", "sort_order"))

    def classes_str(tid: int) -> str:
        ids = teacher_classes.get(tid, set()) | tpl_classes.get(tid, set())
        ids = sorted(ids, key=lambda i: (class_order.get(i, 0), i))
        return "、".join(class_names.get(i, "") for i in ids)

    rows = [
        {
            "tid": t.id,
            "subject": t.subject.name,
            "name": t.name,
            "classes": classes_str(t.id),
        }
        for t in teachers
    ]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_detail_sheet(wb.create_sheet("详细课表"), rows, tpl_count)
    _write_date_sheet(
        wb.create_sheet("周一至周五课"), rows, count, date_kind, COURSE_KINDS, False
    )
    _write_date_sheet(
        wb.create_sheet("周一至周五自习"), rows, count, date_kind, STUDY_KINDS, False
    )
    _write_date_sheet(
        wb.create_sheet("周六日课"), rows, count, date_kind, COURSE_KINDS, True
    )
    _write_date_sheet(
        wb.create_sheet("周六日自习"), rows, count, date_kind, STUDY_KINDS, True
    )
    _write_proctor_sheet(wb.create_sheet("监考"), rows)

    return wb


def _write_date_sheet(ws, rows, count, date_kind, kinds, weekend):
    dates = sorted(
        {d for (d, k) in date_kind if _is_weekend(d) == weekend and k in kinds}
    )

    for col, head in enumerate(["序号", "学科", "姓名", "班级"], start=1):
        ws.cell(row=1, column=col, value=head)
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)

    col = 5
    date_cols: list[tuple[date, list[int]]] = []
    for day in dates:
        first = col
        for kind in kinds:
            ws.cell(row=2, column=col, value=KIND_LABEL[kind])
            col += 1
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=col - 1)
        ws.cell(row=1, column=first, value=_date_label(day))
        date_cols.append((day, list(range(first, col))))
    total_col = col

    for i, row in enumerate(rows):
        r = 3 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=row["subject"])
        ws.cell(row=r, column=3, value=row["name"])
        ws.cell(row=r, column=4, value=row["classes"])
        for day, cols in date_cols:
            for kind, c in zip(kinds, cols):
                v = count.get((row["tid"], day, kind), 0)
                if v:
                    ws.cell(row=r, column=c, value=v)
        if date_cols:
            first_c = date_cols[0][1][0]
            last_c = date_cols[-1][1][-1]
            ws.cell(
                row=r,
                column=total_col,
                value=f"=SUM({get_column_letter(first_c)}{r}:{get_column_letter(last_c)}{r})",
            )
        else:
            ws.cell(row=r, column=total_col, value=0)

    _finish(ws, 2 + len(rows), total_col)


def _write_detail_sheet(ws, rows, tpl_count):
    for col, head in enumerate(["序号", "学科", "姓名", "班级"], start=1):
        ws.cell(row=1, column=col, value=head)
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)

    col = 5
    for weekday in range(1, 8):
        first = col
        for kind in DETAIL_KINDS:
            ws.cell(row=2, column=col, value=KIND_LABEL[kind])
            col += 1
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=col - 1)
        ws.cell(row=1, column=first, value=f"周{WEEKDAY_CN[weekday - 1]}")
    ncols = col - 1

    for i, row in enumerate(rows):
        r = 3 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=row["subject"])
        ws.cell(row=r, column=3, value=row["name"])
        ws.cell(row=r, column=4, value=row["classes"])
        c = 5
        for weekday in range(1, 8):
            for kind in DETAIL_KINDS:
                v = tpl_count.get((row["tid"], weekday, kind), 0)
                if v:
                    ws.cell(row=r, column=c, value=v)
                c += 1

    _finish(ws, 2 + len(rows), ncols)


def _write_proctor_sheet(ws, rows):
    for col, head in enumerate(["序号", "学科", "姓名", "班级"], start=1):
        ws.cell(row=1, column=col, value=head)
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)

    for i, row in enumerate(rows):
        r = 3 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=row["subject"])
        ws.cell(row=r, column=3, value=row["name"])
        ws.cell(row=r, column=4, value=row["classes"])

    _finish(ws, 2 + len(rows), 4, freeze="A3")


# ────────────────────────── 教师个人课表 ──────────────────────────

TEACHING_SLOT_ORDER = [
    "早读",
    "一",
    "二",
    "三",
    "四",
    "五",
    "六",
    "七",
    "八",
    "晚一",
    "晚二",
    "晚三",
]
FLAG_FILL = {
    "swap": "FFE0B2",  # 换课：橙
    "substitute": "BBDEFB",  # 代课：蓝
    "abnormal": "FFCDD2",  # 异常：红
}
FLAG_LABEL = dict(CourseSession.Flag.choices)


def _sheet_title(name: str) -> str:
    for ch in "[]:*?/\\":
        name = name.replace(ch, "_")
    return name[:31]


def build_teacher_workbook(start: date, end: date) -> openpyxl.Workbook:
    slots = list(
        TimeSlot.objects.filter(name__in=TEACHING_SLOT_ORDER, is_active=True)
    )
    slots.sort(key=lambda s: TEACHING_SLOT_ORDER.index(s.name))
    slot_ids = [s.id for s in slots]

    sessions = CourseSession.objects.filter(
        date__range=(start, end), time_slot_id__in=slot_ids
    ).select_related("school_class", "teacher")
    by_teacher: dict[int, dict] = defaultdict(dict)
    for s in sessions:
        if s.teacher_id:
            by_teacher[s.teacher_id][(s.date, s.time_slot_id)] = s

    days: list[date] = []
    cur = start
    while cur <= end:
        days.append(cur)
        cur += timedelta(days=1)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    used: set[str] = set()
    for teacher in Teacher.objects.filter(is_active=True).select_related("subject"):
        title = _sheet_title(f"{teacher.subject.name}-{teacher.name}")
        base = title
        n = 1
        while title in used:
            n += 1
            title = _sheet_title(f"{base}({n})")
        used.add(title)
        ws = wb.create_sheet(title)
        _write_teacher_sheet(ws, teacher, days, slots, by_teacher.get(teacher.id, {}))

    return wb


def _write_teacher_sheet(ws, teacher, days, slots, sess_map):
    headers = ["日期", "星期"] + [s.name for s in slots] + ["合计"]
    for col, head in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=head)
    ncols = len(headers)

    flagged = []
    for i, day in enumerate(days):
        r = 2 + i
        ws.cell(row=r, column=1, value=f"{day.month}.{day.day}")
        ws.cell(row=r, column=2, value=f"周{WEEKDAY_CN[day.isoweekday() - 1]}")
        for j, slot in enumerate(slots):
            session = sess_map.get((day, slot.id))
            if not session:
                continue
            cell = ws.cell(row=r, column=3 + j, value=session.school_class.name)
            if session.flag and session.flag != "normal":
                fill = FLAG_FILL.get(session.flag)
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                flagged.append(
                    (
                        day,
                        slot.name,
                        session.school_class.name,
                        FLAG_LABEL.get(session.flag, ""),
                        session.note,
                    )
                )
        ws.cell(
            row=r,
            column=ncols,
            value=(
                f"=COUNTA({get_column_letter(3)}{r}:"
                f"{get_column_letter(ncols - 1)}{r})"
            ),
        )

    # 空一行后列出所有带标记的条目
    note_start = 2 + len(days) + 1
    ws.cell(row=note_start, column=1, value="标注说明")
    for col, head in enumerate(["日期", "节次", "班级", "标记", "备注"], start=1):
        ws.cell(row=note_start + 1, column=col, value=head)
    for k, (day, slot, cls, flag, note) in enumerate(flagged):
        r = note_start + 2 + k
        ws.cell(row=r, column=1, value=f"{day.month}.{day.day}")
        ws.cell(row=r, column=2, value=slot)
        ws.cell(row=r, column=3, value=cls)
        ws.cell(row=r, column=4, value=flag)
        ws.cell(row=r, column=5, value=note or "")

    grid_rows = 1 + len(days)
    for r in range(1, grid_rows + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER
        ws.cell(row=r, column=ncols).font = Font(bold=True)
    for c in range(1, ncols + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)

    list_end = note_start + 1 + len(flagged)
    for r in range(note_start, list_end + 1):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            cell.alignment = _CENTER
    for c in range(1, 6):
        ws.cell(row=note_start + 1, column=c).font = Font(bold=True)

    ws.row_dimensions[1].height = 22
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 6
    for c in range(3, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10
    ws.freeze_panes = "C2"
